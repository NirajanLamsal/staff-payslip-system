"""Builds a SYNTHETIC (invented) salary workbook for tests. Not real payroll data."""
import sys
from pathlib import Path
import openpyxl
from openpyxl.comments import Comment

HEADERS = ["S.N", "Employe ID", "Name of employees", "Gross Salary", "ABSENT DAYS", "Working Days",
           "Working Day Amount", "Paid leave", "loyalty allowance ", "travelling", "Camp incentive",
           "FIXED SSF DEDECTIONS(11%)", "Total Amount", "Income Tax (1%)", "PayCheque", " Remarks"]
H = {h.strip().lower(): i + 1 for i, h in enumerate(HEADERS)}


def build(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "synthetic 2083"
    ws.append(HEADERS)

    def add(n, id_, name, basic, **kw):
        r = ws.max_row + 1
        ws.cell(r, 1, n); ws.cell(r, 2, id_); ws.cell(r, 3, name); ws.cell(r, 4, basic); ws.cell(r, 6, 30)
        ws.cell(r, H["working day amount"], basic)
        for k, v in kw.items():
            if k == "comment_v":
                ws.cell(r, H["paid leave"]).comment = Comment(v, "x")
            else:
                ws.cell(r, H[k if k in H else k.replace("_", " ")], v)
        return r

    add(1, "T_1", "Ram Many", 25000, loyalty_allowance=2000, travelling=1000, **{"paid_leave": 1000, "comment_v": "Earned Leave: Rs. 500\nRestricted Leave = NPR 500"})
    add(2, "T_2", "Hari Basic", 26000)
    add(3, "T_3", "Sita Travel", 24000, travelling=1000)
    add(4, "T_4", "Bad Comment", 20000, paid_leave=800, comment_v="Earned 500 and 300 leave 200")
    add(5, "T_5", "Empty Comment", 20000, paid_leave=500)
    add(6, "T_6", "Zero Loyalty", 18000, loyalty_allowance=0)
    add(7, "T_7", "A Very Long Employee Name With Many Parts Sharma Shrestha Bhattarai Gurung", 30000, camp_incentive=500,
        **{"fixed ssf dedections(11%)": 3000})
    r = add(8, "T_8", "Formula Person", 12000)
    ws.cell(r, H["working day amount"], f"=D{r}-(D{r}/F{r})*5")
    ws.cell(r, H["total amount"], f"=G{r}")
    ws.cell(r, H["income tax (1%)"], f"=M{r}*1%")
    ws.cell(r, H["paycheque"], f"=M{r}-N{r}")
    add(9, "T_9", "Text Amount", 10000, travelling="abc")
    add(10, "T_9", "Duplicate Id", 9000)
    add(11, "T_11", None, 9000)
    ws2 = wb.create_sheet("Other notes")
    ws2["A1"] = "not a salary sheet"
    wb.save(path)


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).parent.parent / "sample_data" / "SYNTHETIC_TEST_DATA.xlsx"
    build(out)
    print(out)
