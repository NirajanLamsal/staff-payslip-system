"""Reading the salary workbook (read-only; the uploaded file is never modified or saved)."""
from __future__ import annotations

import io
import re
import zipfile
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils import get_column_letter

from . import config as cfg
from .comments import clean_comment
from .models import Cell, ColumnInfo, Flag, INFO, RawRow, SheetData, WARNING
from .recalc import libreoffice_path, recalculate
from .util import UserError, clean_text, norm_header

SUPPORTED = (".xlsx", ".xlsm")
NAME_ROLE_KEYS = {"name of employees", "name of employee", "employee name", "name", "staff name",
                  "name of staff", "employees name"}


MAX_UNZIPPED_BYTES = 150 * 1024 * 1024   # a payroll workbook is a few hundred KB; this stops zip bombs
MAX_ZIP_ENTRIES = 2000


def _check_archive(data: bytes) -> None:
    """Reject workbooks that are not a sane .xlsx container (decompression bombs, absurd entry counts, odd paths)."""
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as zf:
            infos = zf.infolist()
            if len(infos) > MAX_ZIP_ENTRIES or sum(i.file_size for i in infos) > MAX_UNZIPPED_BYTES:
                raise UserError("This workbook is unusually large inside and was not opened for safety. "
                                "Open it in Excel, save a clean copy and upload that.")
            for i in infos:
                if i.filename.startswith(("/", "\\")) or ".." in i.filename.replace("\\", "/").split("/"):
                    raise UserError("This file has an unsafe internal structure and was not opened.")
            if not any(i.filename == "[Content_Types].xml" for i in infos):
                raise zipfile.BadZipFile("not an OOXML package")
    except zipfile.BadZipFile:
        raise UserError("This file could not be opened as an Excel workbook. It may be damaged, "
                        "password protected or not a real .xlsx file. Please open it in Excel or "
                        "LibreOffice, save it again as .xlsx and upload that copy.")


def _open(data: bytes, data_only: bool):
    _check_archive(data)
    try:
        return openpyxl.load_workbook(io.BytesIO(data), data_only=data_only)
    except (zipfile.BadZipFile, KeyError, OSError, ValueError, openpyxl.utils.exceptions.InvalidFileException):
        raise UserError("This file could not be opened as an Excel workbook. It may be damaged, "
                        "password protected or not a real .xlsx file. Please open it in Excel or "
                        "LibreOffice, save it again as .xlsx and upload that copy.")


def check_filename(filename: str) -> None:
    low = filename.lower()
    if not low.endswith(SUPPORTED):
        if low.endswith(".xls"):
            raise UserError("Old .xls files are not supported. Open it in Excel or LibreOffice and "
                            "use 'Save As' to create an .xlsx copy.")
        raise UserError("Unsupported file type. Please upload an Excel workbook (.xlsx).")


def sheet_names(data: bytes, filename: str) -> List[str]:
    check_filename(filename)
    wb = _open(data, data_only=False)
    return list(wb.sheetnames)


def guess_sheet(data: bytes, settings: Dict[str, Any]) -> str:
    """The sheet whose first rows look most like an employee table."""
    wb = _open(data, data_only=True)
    roles = cfg.effective_roles(settings)
    best, best_score = wb.sheetnames[0], -1
    for ws in wb.worksheets:
        if ws.sheet_state != "visible":
            continue
        score = 0
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 30), values_only=True):
            known = sum(1 for v in row if v is not None and norm_header(v) in roles)
            score = max(score, known)
        if score > best_score:
            best, best_score = ws.title, score
    return best


def _find_header_row(ws, roles: Dict[str, str]) -> Optional[int]:
    best_row, best_score = None, 0
    for r in range(1, min(ws.max_row, 40) + 1):
        keys = [norm_header(c.value) for c in ws[r] if c.value is not None]
        if not any(k in NAME_ROLE_KEYS or roles.get(k) == cfg.ROLE_NAME for k in keys):
            continue
        score = sum(1 for k in keys if k in roles)
        if score > best_score:
            best_row, best_score = r, score
    return best_row


def _cell_text(value: Any) -> str:
    return clean_text(value)[:120]


def read_sheet(data: bytes, filename: str, sheet_name: Optional[str], settings: Dict[str, Any],
               header_row: Optional[int] = None, allow_recalc: bool = True) -> SheetData:
    check_filename(filename)
    wb_f = _open(data, data_only=False)
    names = list(wb_f.sheetnames)
    if sheet_name is None:
        sheet_name = guess_sheet(data, settings)
    if sheet_name not in names:
        raise UserError(f"The worksheet '{sheet_name}' was not found in this workbook.")
    wb_v = _open(data, data_only=True)
    ws_f, ws_v = wb_f[sheet_name], wb_v[sheet_name]
    flags: List[Flag] = []
    recalculated = False

    # --- formulas without a stored result (file produced by a program, or never calculated)
    def count_uncached(wsf, wsv) -> int:
        n = 0
        for row in wsf.iter_rows():
            for c in row:
                v = c.value
                is_formula = (isinstance(v, str) and v.startswith("=")) or hasattr(v, "text")
                if is_formula and wsv[c.coordinate].value is None:
                    n += 1
        return n

    missing = count_uncached(ws_f, ws_v)
    if missing and allow_recalc and libreoffice_path():
        fresh = recalculate(data)
        if fresh:
            try:
                wb_v2 = _open(fresh, data_only=True)
                if sheet_name in wb_v2.sheetnames:
                    ws_v2 = wb_v2[sheet_name]
                    if count_uncached(ws_f, ws_v2) < missing:
                        ws_v, recalculated = ws_v2, True
                        flags.append(Flag(INFO, f"{missing} formula cells had no stored result; "
                                                "they were recalculated on this computer with LibreOffice."))
            except UserError:
                pass

    roles = cfg.effective_roles(settings)
    if header_row is None:
        header_row = _find_header_row(ws_v, roles)
    if header_row is None:
        raise UserError(f"Could not find the header row with employee names on the sheet '{sheet_name}'. "
                        "Check that you selected the right worksheet, or enter the header row number "
                        "manually.")

    if ws_f.merged_cells.ranges:
        flags.append(Flag(INFO, f"The sheet has {len(ws_f.merged_cells.ranges)} merged cell range(s); "
                                "values are read from the top-left cell of each."))

    # --- columns
    columns: List[ColumnInfo] = []
    seen_keys: Dict[str, int] = {}
    max_col = ws_v.max_column
    for c in range(1, max_col + 1):
        raw = ws_v.cell(header_row, c).value
        letter = get_column_letter(c)
        has_data = any(ws_v.cell(r, c).value is not None for r in range(header_row + 1, ws_v.max_row + 1))
        if raw is None or clean_text(raw) == "":
            if has_data:
                header = f"(no header, column {letter})"
                key = norm_header(header)
                columns.append(ColumnInfo(c, letter, header, key, cfg.ROLE_UNMAPPED, header))
            continue
        header = clean_text(raw)
        key = norm_header(raw)
        if key in seen_keys:
            flags.append(Flag(WARNING, f"The header '{header}' appears more than once (columns "
                                       f"{get_column_letter(seen_keys[key])} and {letter})."))
            key = f"{key}#{letter}"
        seen_keys[key] = c
        role = roles.get(key) or cfg.guess_role(key)
        columns.append(ColumnInfo(c, letter, header, key, role, cfg.display_label(header, settings)))

    name_cols = [c for c in columns if c.role == cfg.ROLE_NAME]
    if not name_cols:
        raise UserError("No 'employee name' column was found. Use 'Column mapping' to tell the app "
                        "which column holds the employee names.")
    id_cols = [c for c in columns if c.role == cfg.ROLE_ID]

    # --- rows
    rows: List[RawRow] = []
    stray: List[Tuple[int, str, str]] = []
    for r in range(header_row + 1, ws_v.max_row + 1):
        cells: Dict[int, Cell] = {}
        any_value = False
        for col in columns:
            cv = ws_v.cell(r, col.index)
            cf = ws_f.cell(r, col.index)
            val = cv.value
            formula = None
            fv = cf.value
            if isinstance(fv, str) and fv.startswith("="):
                formula = fv
            elif hasattr(fv, "text"):
                formula = "=" + str(getattr(fv, "text", ""))
            comment = clean_comment(cf.comment.text) if cf.comment else ""
            is_err = cv.data_type == "e"
            if isinstance(val, str) and not val.strip():
                val = None
            if val is not None or comment:
                any_value = True
            cells[col.index] = Cell(value=val, formula=formula, comment=comment, is_error=is_err)
        if not any_value:
            continue
        name_val = next((cells[c.index].value for c in name_cols if cells[c.index].value is not None), None)
        id_val = next((cells[c.index].value for c in id_cols if cells[c.index].value is not None), None)
        if name_val is None and id_val is None:
            for col in columns:
                cl = cells[col.index]
                if cl.value is not None:
                    stray.append((r, col.letter, _cell_text(cl.value)))
            continue
        if name_val is not None and re.match(r"^\s*(grand\s+)?total\b", str(name_val), re.I):
            flags.append(Flag(INFO, f"Row {r} ('{_cell_text(name_val)}') looks like a totals row and was skipped."))
            continue
        rows.append(RawRow(row=r, cells=cells))

    for (r, letter, text) in stray:
        flags.append(Flag(INFO, f"Cell {letter}{r} contains text outside the employee rows and was not used: '{text}'"))

    if not rows:
        raise UserError(f"No employee rows were found below the header on the sheet '{sheet_name}'.")

    return SheetData(sheet_name=sheet_name, sheet_names=names, header_row=header_row, columns=columns,
                     rows=rows, stray=stray, global_flags=flags, recalculated=recalculated)
